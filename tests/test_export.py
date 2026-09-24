"""R4.2: Excel exports — sheets, styled header, filters, timestamped file name."""

import io
import re

from openpyxl import load_workbook

from conftest import REG_INPUT


def _save_arena(client, user):
    arena = client.post("/regression/arena", json=REG_INPUT).json()
    items = [{"task": "regression", "model": r["model"], "input": arena["input"],
              "predicted_value": r["prediction"], "runtime_ms": r["runtime_ms"], "actual_value": 1.3}
             for r in arena["results"]]
    client.post("/db/predictions", json={"items": items}, headers=user["headers"])


def _filename(res) -> str:
    return re.search(r'filename="([^"]+)"', res.headers["content-disposition"]).group(1)


def test_export_history(client, user):
    _save_arena(client, user)
    res = client.get("/db/export/predictions.xlsx", headers=user["headers"])
    assert res.status_code == 200
    assert res.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert re.fullmatch(r"lich_su_du_doan_\d{8}_\d{6}\.xlsx", _filename(res))

    wb = load_workbook(io.BytesIO(res.content))
    assert wb.sheetnames == ["Lịch sử dự đoán", "Bộ lọc"]
    ws = wb["Lịch sử dự đoán"]
    assert ws.max_row == 1 + 5
    header = ws["A1"]
    assert header.font.bold and header.fill.fgColor.rgb.endswith("6C5CE7")
    assert ws.column_dimensions["E"].width > 10  # auto-sized input column


def test_export_history_filtered_by_model(client, user):
    _save_arena(client, user)
    res = client.get("/db/export/predictions.xlsx", params={"model": "ridge"}, headers=user["headers"])
    ws = load_workbook(io.BytesIO(res.content))["Lịch sử dự đoán"]
    assert ws.max_row == 2 and ws["D2"].value == "ridge"


def test_export_requires_login(client):
    assert client.get("/db/export/predictions.xlsx").status_code == 401


def test_export_model_comparison(client, user):
    metrics = client.get("/regression/metrics").json()
    client.post("/db/model-runs", headers=user["headers"], json={
        "best_model": metrics["best_model"], "models": metrics["models"],
        "train_size": metrics["data"]["train_size"], "test_size": metrics["data"]["test_size"],
    })
    res = client.get("/db/export/model-runs.xlsx", headers=user["headers"])
    assert res.status_code == 200
    assert re.fullmatch(r"so_sanh_mo_hinh_\d{8}_\d{6}\.xlsx", _filename(res))
    wb = load_workbook(io.BytesIO(res.content))
    assert wb.sheetnames == ["So sánh mô hình", "Lịch sử train"]
    ws = wb["So sánh mô hình"]
    labels = [ws.cell(row=r, column=1).value for r in range(2, 7)]
    assert labels == [m["label"] for m in metrics["models"]]
    assert ws["A1"].font.bold
